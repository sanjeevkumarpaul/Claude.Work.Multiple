#!/usr/bin/env python3
"""Exact-layout OCR: image -> monospaced text (Courier New) preserving spacing.

Usage:  python exact_ocr.py image.png [-o out_basename] [--lang eng] [--psm 6]
Writes  <out>.txt  (plain, fixed-width layout)
        <out>.html (same text rendered in Courier New)
        <out>.docx (Courier New, only if python-docx is installed)
"""
import argparse
import html
import statistics
import sys
from pathlib import Path

import pytesseract
from PIL import Image, ImageOps, ImageStat

MIN_HEIGHT = 1500  # upscale small images so Tesseract sees enough pixels


def load(path):
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    if img.height < MIN_HEIGHT:
        s = MIN_HEIGHT / img.height
        img = img.resize((round(img.width * s), MIN_HEIGHT), Image.LANCZOS)
    gray = ImageOps.autocontrast(img.convert("L"))
    if ImageStat.Stat(gray).mean[0] < 110:  # light-on-dark
        gray = ImageOps.invert(gray)
    return gray


def read_words(img, lang, psm):
    d = pytesseract.image_to_data(
        img, lang=lang, config=f"--oem 1 --psm {psm} -c preserve_interword_spaces=1",
        output_type=pytesseract.Output.DICT)
    words = []
    for i, t in enumerate(d["text"]):
        if t.strip() and float(d["conf"][i]) >= 0:
            words.append(dict(t=t.strip(), x=d["left"][i], y=d["top"][i],
                              w=d["width"][i], h=d["height"][i]))
    return words


def to_layout(words):
    if not words:
        return ""
    # character cell width: median of word pixel width / chars (longer words are more reliable)
    ws = [w["w"] / len(w["t"]) for w in words if len(w["t"]) >= 3] or \
         [w["w"] / len(w["t"]) for w in words]
    cw = statistics.median(ws)
    lh = statistics.median(w["h"] for w in words)

    # cluster words into visual rows by vertical centre
    words.sort(key=lambda w: w["y"] + w["h"] / 2)
    rows, cur = [], [words[0]]
    for w in words[1:]:
        cy = statistics.fmean(c["y"] + c["h"] / 2 for c in cur)
        if abs(w["y"] + w["h"] / 2 - cy) < 0.6 * lh:
            cur.append(w)
        else:
            rows.append(cur)
            cur = [w]
    rows.append(cur)

    cys = [statistics.fmean(w["y"] + w["h"] / 2 for w in r) for r in rows]
    gaps = [b - a for a, b in zip(cys, cys[1:])]
    pitch = max(min(gaps), lh) if gaps else lh  # one text line of vertical space
    x0 = min(w["x"] for w in words)
    out, prev_cy = [], None
    for row in rows:
        cy = statistics.fmean(w["y"] + w["h"] / 2 for w in row)
        if prev_cy is not None:  # keep vertical gaps as blank lines
            out.extend([""] * max(0, round((cy - prev_cy) / pitch) - 1))
        prev_cy = cy
        line = ""
        for w in sorted(row, key=lambda w: w["x"]):
            col = round((w["x"] - x0) / cw)
            if col <= len(line):  # collision/tight: single space separator
                col = len(line) + (1 if line else 0)
            line += " " * (col - len(line)) + w["t"]
        out.append(line.rstrip())
    return "\n".join(out) + "\n"


def write_outputs(text, base):
    base.with_suffix(".txt").write_text(text, encoding="utf-8")
    base.with_suffix(".html").write_text(
        '<!doctype html><meta charset="utf-8"><title>OCR</title>'
        '<pre style="font-family:\'Courier New\',Courier,monospace;font-size:14px">'
        f"{html.escape(text)}</pre>", encoding="utf-8")
    written = [".txt", ".html"]
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Courier New", Pt(10)
        for ln in text.rstrip("\n").split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(ln).font.name = "Courier New"
        doc.save(base.with_suffix(".docx"))
        written.append(".docx")
    except ImportError:
        pass
    return [str(base.with_suffix(e)) for e in written]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("image")
    ap.add_argument("-o", "--out", help="output basename (default: image name)")
    ap.add_argument("--lang", default="eng")
    ap.add_argument("--psm", type=int, default=6, help="Tesseract page mode (6=block, 4=columns, 11=sparse)")
    a = ap.parse_args()
    text = to_layout(read_words(load(a.image), a.lang, a.psm))
    sys.stdout.write(text)
    base = Path(a.out) if a.out else Path(a.image).with_suffix("")
    print("\nSaved:", ", ".join(write_outputs(text, base)), file=sys.stderr)


if __name__ == "__main__":
    main()
